# -*- coding: utf-8 -*-
"""
Dựng bản PowerPoint (.pptx) cho pitch HAWEE, phản chiếu hawee-pitch-deck.html (v3.1).
Native, chỉnh sửa được; biểu đồ là chart thật trong PowerPoint.
Chạy: .venv-pptx/bin/python hawee-pitch/build_hawee_pptx.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

NAVY=RGBColor(0x0F,0x27,0x47); NAVY2=RGBColor(0x1D,0x3F,0x63)
ACCENT=RGBColor(0x1C,0x6F,0xB8); ACCENT2=RGBColor(0x3F,0x9B,0xD6)
GOLD=RGBColor(0xB7,0x89,0x2F); GREEN=RGBColor(0x1F,0x9D,0x6B)
AMBER=RGBColor(0xC9,0x8A,0x2E); VIOLET=RGBColor(0x6B,0x54,0xB8)
TEAL=RGBColor(0x1F,0x8A,0x8A)
SOFT=RGBColor(0xF5,0xF7,0xFB); LINE=RGBColor(0xE4,0xE9,0xF0)
MUTED=RGBColor(0x5D,0x6B,0x7E); INK=RGBColor(0x15,0x27,0x3D); WHITE=RGBColor(0xFF,0xFF,0xFF)

prs=Presentation()
prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]
SW,SH=prs.slide_width,prs.slide_height

def slide():
    return prs.slides.add_slide(BLANK)

def bg(s,color):
    s.background.fill.solid(); s.background.fill.fore_color.rgb=color

def box(s,x,y,w,h,fill=None,linec=None,round_=True,shadow=False):
    shp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,x,y,w,h)
    shp.shadow.inherit=False
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb=fill
    if linec is None: shp.line.fill.background()
    else: shp.line.color.rgb=linec; shp.line.width=Pt(0.75)
    return shp

def txt(s,x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,sp=1.0):
    tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    if isinstance(runs[0],tuple): runs=[runs]
    for i,para in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; p.line_spacing=sp; p.space_after=Pt(3)
        for (t,sz,c,b) in para:
            r=p.add_run(); r.text=t; f=r.font
            f.size=Pt(sz); f.color.rgb=c; f.bold=b; f.name="Calibri"
    return tb

def kicker(s,t,color=ACCENT):
    txt(s,Inches(0.7),Inches(0.42),Inches(11.9),Inches(0.35),[[(t.upper(),12,color,True)]])

def title(s,t,em=None,emc=ACCENT):
    runs=[(t,30,NAVY,True)]
    if em: runs.append((" "+em,30,emc,True))
    txt(s,Inches(0.7),Inches(0.78),Inches(12),Inches(1.0),[runs])

def src(s,t):
    txt(s,Inches(0.7),Inches(7.0),Inches(12),Inches(0.35),[[(t,9,RGBColor(0x9A,0xA7,0xB6),False)]])

def quote(s,t,y=Inches(5.7),w=Inches(11.9)):
    bar=box(s,Inches(0.7),y,Inches(0.06),Inches(0.9),GOLD)
    txt(s,Inches(0.95),y,w,Inches(0.9),[[(t,15,NAVY,True)]],anchor=MSO_ANCHOR.MIDDLE,sp=1.1)

def bullets(s,items,x=Inches(0.7),y=Inches(1.85),w=Inches(7.4),sz=15,gap=True):
    tb=s.shapes.add_textbox(x,y,w,Inches(4.6)); tf=tb.text_frame; tf.word_wrap=True
    tf.margin_left=0;tf.margin_top=0
    for i,it in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.space_after=Pt(10 if gap else 4); p.line_spacing=1.15
        rb=p.add_run(); rb.text="▪  "; rb.font.size=Pt(sz); rb.font.color.rgb=ACCENT; rb.font.bold=True
        if isinstance(it,str): it=[(it,sz,INK,False)]
        for (t,zz,c,b) in it:
            r=p.add_run(); r.text=t; r.font.size=Pt(zz); r.font.color.rgb=c; r.font.bold=b; r.font.name="Calibri"
    return tb

def cards(s,data,y=Inches(2.0),h=Inches(2.5),cols=None):
    n=len(data); cols=cols or n
    gap=Inches(0.25); total=SW-Inches(1.4)
    cw=int((total-gap*(cols-1))/cols)
    for i,(tag,head,body) in enumerate(data):
        x=Inches(0.7)+i*(cw+gap)
        c=box(s,x,y,cw,h,SOFT,LINE)
        txt(s,x+Inches(0.22),y+Inches(0.2),Emu(cw)-Inches(0.4),Inches(0.3),[[(tag.upper(),10,ACCENT,True)]])
        txt(s,x+Inches(0.22),y+Inches(0.52),Emu(cw)-Inches(0.4),Inches(0.5),[[(head,16,NAVY,True)]])
        txt(s,x+Inches(0.22),y+Inches(1.05),Emu(cw)-Inches(0.4),Emu(h)-Inches(1.2),[[(body,12,MUTED,False)]],sp=1.1)

def kpis(s,data,y=Inches(2.3)):
    n=len(data); gap=Inches(0.4); total=SW-Inches(1.4)
    cw=int((total-gap*(n-1))/n)
    for i,(v,l,c) in enumerate(data):
        x=Inches(0.7)+i*(cw+gap)
        txt(s,x,y,cw,Inches(0.9),[[(v,40,c,True)]])
        txt(s,x,y+Inches(0.95),cw,Inches(0.8),[[(l,12,MUTED,False)]],sp=1.05)

def flow(s,steps,y=Inches(2.6),h=Inches(1.9)):
    n=len(steps); aw=Inches(0.45); gap=Inches(0.12)
    total=SW-Inches(1.4)
    cw=int((total-(n-1)*(aw+2*gap))/n)
    x=Inches(0.7)
    for i,(si,head,body,fill) in enumerate(steps):
        c=box(s,x,y,cw,h,fill,LINE)
        txt(s,x+Inches(0.18),y+Inches(0.16),Emu(cw)-Inches(0.36),Inches(0.3),[[(si.upper(),10,ACCENT,True)]])
        txt(s,x+Inches(0.18),y+Inches(0.5),Emu(cw)-Inches(0.36),Inches(0.6),[[(head,13,NAVY,True)]],sp=1.0)
        txt(s,x+Inches(0.18),y+Inches(1.12),Emu(cw)-Inches(0.36),Inches(0.7),[[(body,10.5,MUTED,False)]],sp=1.0)
        x=Emu(x)+cw
        if i<n-1:
            txt(s,x+gap,y,aw,h,[[("›",24,GOLD,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
            x=Emu(x)+aw+2*gap

def two_panel(s,left,right,y=Inches(1.95),h=Inches(3.5)):
    gap=Inches(0.55); pw=int((SW-Inches(1.4)-Inches(0.45)-gap*2)/2)
    arrow=Inches(0.45)
    def panel(x,hd,items,fill,hc):
        box(s,x,y,pw,h,fill,LINE)
        txt(s,x+Inches(0.25),y+Inches(0.2),Emu(pw)-Inches(0.5),Inches(0.35),[[(hd.upper(),11,hc,True)]])
        tb=s.shapes.add_textbox(x+Inches(0.25),y+Inches(0.62),Emu(pw)-Inches(0.5),Emu(h)-Inches(0.8))
        tf=tb.text_frame; tf.word_wrap=True; tf.margin_left=0;tf.margin_top=0
        for i,it in enumerate(items):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.space_after=Pt(8); p.line_spacing=1.1
            rb=p.add_run(); rb.text="▪  "; rb.font.size=Pt(12); rb.font.color.rgb=hc; rb.font.bold=True
            r=p.add_run(); r.text=it; r.font.size=Pt(12.5); r.font.color.rgb=INK
    panel(Inches(0.7),left[0],left[1],RGBColor(0xFB,0xF6,0xEE),AMBER)
    txt(s,Inches(0.7)+pw+gap,y,arrow,h,[[("›",26,GOLD,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    panel(Inches(0.7)+pw+gap+arrow+gap,right[0],right[1],RGBColor(0xEE,0xF8,0xF3),GREEN)

def table(s,headers,rows,y=Inches(1.9),x=Inches(0.7),w=None,fs=11):
    w=w or (SW-Inches(1.4))
    nr=len(rows)+1; nc=len(headers)
    h=Inches(0.42)*nr
    gf=s.shapes.add_table(nr,nc,x,y,w,h); t=gf.table
    def fill_cell(cell,val,color,bold):
        cell.vertical_anchor=MSO_ANCHOR.MIDDLE
        tf2=cell.text_frame; tf2.clear(); p=tf2.paragraphs[0]
        r=p.add_run(); r.text=str(val) if val not in (None,"") else " "
        r.font.size=Pt(fs); r.font.bold=bold; r.font.color.rgb=color; r.font.name="Calibri"
    for j,hd in enumerate(headers):
        cell=t.cell(0,j); cell.fill.solid(); cell.fill.fore_color.rgb=NAVY
        fill_cell(cell,hd,WHITE,True)
    for i,row in enumerate(rows):
        for j,val in enumerate(row):
            cell=t.cell(i+1,j); cell.fill.solid()
            cell.fill.fore_color.rgb=WHITE if i%2==0 else SOFT
            fill_cell(cell,val,INK,False)
    return gf

def add_chart(s,ctype,cats,series,x,y,w,h,title_=None,legend=True):
    cd=CategoryChartData(); cd.categories=cats
    for nm,vals in series: cd.add_series(nm,vals)
    gf=s.shapes.add_chart(ctype,x,y,w,h,cd); ch=gf.chart
    ch.has_title=bool(title_)
    if title_:
        ch.chart_title.text_frame.text=title_
        ch.chart_title.text_frame.paragraphs[0].runs[0].font.size=Pt(11)
    ch.has_legend=legend
    if legend:
        ch.legend.position=XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout=False
        ch.legend.font.size=Pt(9)
    try:
        for plot in ch.plots: plot.has_data_labels=False
    except Exception: pass
    return ch

def cover(s,kick,big,sub,foot):
    bg(s,NAVY)
    box(s,0,0,SW,SH,NAVY2).fill.fore_color.rgb=NAVY2
    s.shapes[-1].fill.background()
    txt(s,Inches(0.9),Inches(1.5),Inches(11),Inches(0.4),[[(kick.upper(),13,RGBColor(0x9E,0xC7,0xE8),True)]])
    txt(s,Inches(0.9),Inches(2.0),Inches(11.5),Inches(2.4),[[(line,46,WHITE,True)] for line in big.split("\n")],sp=1.05)
    txt(s,Inches(0.9),Inches(4.5),Inches(10.5),Inches(1.0),[[(sub,17,RGBColor(0xCF,0xDD,0xED),False)]],sp=1.2)
    box(s,Inches(0.9),Inches(6.1),Inches(11.5),Pt(1),RGBColor(0x39,0x55,0x73))
    txt(s,Inches(0.9),Inches(6.3),Inches(11.5),Inches(0.5),[[(foot,12,RGBColor(0x9F,0xB6,0xCE),False)]])

def divider(s,part,big,sub):
    bg(s,NAVY)
    txt(s,Inches(0.9),Inches(2.4),Inches(11),Inches(0.4),[[(part.upper(),13,RGBColor(0x9E,0xC7,0xE8),True)]])
    txt(s,Inches(0.9),Inches(2.9),Inches(11.5),Inches(1.3),[[(big,40,WHITE,True)]])
    txt(s,Inches(0.9),Inches(4.3),Inches(10.5),Inches(0.9),[[(sub,16,RGBColor(0xCF,0xDD,0xED),False)]],sp=1.2)

def head(s,kick,t,em=None,kc=ACCENT,emc=ACCENT):
    bg(s,WHITE); kicker(s,kick,kc); title(s,t,em,emc)

# ============ SLIDES ============

# 1 Cover
s=slide(); cover(s,"Đề xuất chiến lược · Nhiệm kỳ III 2025 đến 2030",
    "HAWEE: Từ Mạng lưới\nđến Nền tảng",
    "Một nền tảng số xây riêng cho HAWEE, hiện thực hóa tinh thần “Hội tụ, Kết nối, Vươn tầm”.",
    "HAWEE      Bản trình bày giải pháp      2026")

# 2 Context
s=slide(); head(s,"Bối cảnh nhiệm kỳ III","HAWEE đang ở","thời điểm vàng để bứt phá")
kpis(s,[("500+","hội viên nữ lãnh đạo",NAVY),("4 hạng","Cơ Bản đến Lan Tỏa",ACCENT),
        ("45","thành viên Ban Chấp hành",NAVY),("10 năm","hành trình 2015 đến 2025",GOLD)],y=Inches(1.9))
bullets(s,["Định hướng mới: thúc đẩy giao thương, mở rộng thị trường, kết nối đối tác cho hội viên",
           "Mở rộng HAWEE Mentoring theo chiều sâu ngành nghề"],y=Inches(3.7),w=Inches(11.9))
quote(s,"Tầm nhìn đã rõ. Việc còn lại là chọn công cụ vận hành xứng tầm với 500+ hội viên.")
src(s,"Nguồn: hawee.vn; Đại hội nhiệm kỳ III (theleader.vn, phunuvietnam.vn) 03/2025")

# 3 Market
s=slide(); head(s,"Cơ hội thị trường","HAWEE đứng trên một","làn sóng đang lên",kc=GREEN)
add_chart(s,XL_CHART_TYPE.COLUMN_CLUSTERED,["2015 (~21%)","2025 (~25%)","2030 (30%)"],
          [("DN do nữ làm chủ",[21,25,30])],Inches(0.7),Inches(1.9),Inches(5.6),Inches(4.4),
          "Tỷ lệ doanh nghiệp Việt do nữ làm chủ",legend=False)
bullets(s,["~25% doanh nghiệp Việt do nữ làm chủ, hướng 30% vào 2030",
           "97% doanh nghiệp do nữ làm chủ còn nhỏ, rất cần kết nối và cố vấn",
           "Mạng lưới nữ doanh nhân Việt Nam thuộc nhóm hiệu quả nhất ASEAN"],
        x=Inches(6.6),y=Inches(2.1),w=Inches(6.0))
quote(s,"HAWEE có vị thế dẫn dắt. Nền tảng số biến mạng lưới thành giá trị nhìn thấy được.",y=Inches(5.5),w=Inches(6.0))
src(s,"Nguồn: VCCI, WISE Vietnam, VietnamPlus (2025). Cột 2015 mang tính minh họa xu hướng.")

# 4 Single solution thesis
s=slide(); head(s,"Định hướng giải pháp","Một nền tảng","xây riêng cho HAWEE")
cards(s,[("Của HAWEE","Sở hữu trọn vẹn","Dữ liệu và sản phẩm thuộc về Hội, không phụ thuộc bên thứ ba."),
         ("Đúng bản sắc","May đo theo Hội","Thiết kế đúng 4 hạng, chương trình Mentoring và định hướng giao thương."),
         ("Sát người Việt","Gắn liền Zalo","Nhắc nhở, chăm sóc qua Zalo; thanh toán nội địa quen thuộc.")],y=Inches(2.1),h=Inches(2.6))
quote(s,"Đây không phải phần mềm thuê sẵn. Đây là tài sản số riêng của HAWEE, lớn lên cùng Hội qua từng nhiệm kỳ.")

# 5 Divider I
s=slide(); divider(s,"Phần I","Điểm nghẽn và cơ hội",
    "Nhìn thẳng vào nút thắt hôm nay, để thấy rõ dư địa tăng trưởng ngày mai.")

# 6 Pain map 7+3
s=slide(); head(s,"Bản đồ điểm nghẽn","Không chỉ 7, mà là","7 cốt lõi và 3 mở rộng")
table(s,["7 điểm nghẽn cốt lõi","3 điểm nghẽn mở rộng"],
      [["01 Dữ liệu phân tán","08 Hội viên mới ít gắn kết 90 ngày đầu"],
       ["02 Gia hạn thủ công","09 Chưa ghi nhận hội viên năng động"],
       ["03 Tương tác theo nhịp sự kiện","10 Chi hội & truyền thông phân tán"],
       ["04 Chưa đo được giá trị",""],
       ["05 Giao thương rời rạc",""],
       ["06 Mentoring thủ công",""],
       ["07 Hồ sơ hội viên sơ sài",""]],y=Inches(1.9),w=Inches(11.0),fs=12)
src(s,"Điểm nghẽn 08 đến 10 bổ sung từ nghiên cứu thị trường và đặc thù cơ cấu HAWEE.")

# 7 Growth room
s=slide(); head(s,"Dư địa đang chờ khai mở","Mỗi hội viên ở lại là","một giá trị nhân lên",kc=GREEN)
flow(s,[("Hôm nay","500 hội viên","Cơ sở hội viên hiện tại",SOFT),
        ("Khi có nền tảng","Giữ chân vững hơn","Nhắc & chăm sóc tự động",RGBColor(0xEE,0xF8,0xF3)),
        ("Kết quả","Mạng lưới mạnh thêm","Giá trị giao thương lớn dần",RGBColor(0xEE,0xF6,0xFB))],y=Inches(1.95),h=Inches(1.8))
kpis(s,[("+12 điểm %","tỷ lệ hội viên tiếp tục đồng hành",GREEN),
        ("+1,2 tỷ","VND/năm hội phí giữ lại (minh họa)",NAVY),
        ("Bền vững","mạng lưới & uy tín tích lũy",ACCENT)],y=Inches(4.1))
src(s,"Mô hình minh họa: giả định 500 hội viên, hội phí 20 triệu/người. Thay số thực khi khảo sát.")

# 8 Proven model
s=slide(); head(s,"Vững tâm khi quyết định","Mô hình đã được","kiểm chứng ngay tại Việt Nam",kc=GREEN)
add_chart(s,XL_CHART_TYPE.DOUGHNUT,["Tiếp tục đồng hành","Còn lại"],
          [("Giữ chân",[82,18])],Inches(0.7),Inches(1.9),Inches(4.6),Inches(4.2),
          "Tỷ lệ hội viên tiếp tục đồng hành (82%)",legend=True)
bullets(s,["Trung bình tổ chức có nền tảng đạt ~82%, so với khoảng 55% khi làm thủ công",
           "EuroCham Việt Nam (~1.400 hội viên) đã vận hành trên nền tảng quản lý hội viên",
           "Hàng nghìn hội & hiệp hội trên thế giới dùng mô hình tương tự"],
        x=Inches(5.7),y=Inches(2.1),w=Inches(6.9))
quote(s,"Mô hình đã được chứng minh. HAWEE chỉ cần một bản may đo theo bản sắc của mình.",y=Inches(5.4),w=Inches(6.9))
src(s,"Nguồn: tham chiếu nền tảng quản lý hội viên (Glue Up); EuroCham VN. Mức 55% minh họa.")

# 9 Solution 3 layers
s=slide(); head(s,"Giải pháp tổng quan","Một nền tảng,","ba lớp giá trị")
cards(s,[("Lớp 1 · Nền móng","Dữ liệu & Hội viên","Kho dữ liệu tập trung · Hồ sơ 360° · 4 hạng & gia hạn tự động · Nhắc qua Zalo & email"),
         ("Lớp 2 · Giá trị","Gắn kết & Giao thương","Sự kiện, vé & quét QR · Sàn giao thương · Mentoring · Tích điểm & vinh danh"),
         ("Lớp 3 · Vươn tầm","Thông minh & Mở rộng","Bảng điều hành sức khỏe · Chăm sóc sớm · Trợ lý ảo · Hiệu quả tài trợ")],
      y=Inches(2.0),h=Inches(3.0))
quote(s,"Một nơi duy nhất cho hội viên, Ban Chấp hành và đối tác.")

# 10 Map issue->solution
s=slide(); head(s,"Bản đồ điểm nghẽn đến giải pháp","Mỗi nút thắt,","một lời giải rõ ràng")
table(s,["#","Điểm nghẽn","Giải pháp","Kết quả nhìn thấy"],
      [["1","Dữ liệu phân tán","Kho dữ liệu tập trung & Hồ sơ 360°","Một nguồn duy nhất"],
       ["2","Gia hạn thủ công","Quản lý hạng & gia hạn tự động","Giữ chân vững"],
       ["3","Tương tác theo nhịp","Sự kiện, quét QR & bảng tin cộng đồng","Gắn kết liên tục"],
       ["4","Chưa đo được giá trị","Bảng điều hành sức khỏe Hội","Quyết định bằng số"],
       ["5","Giao thương rời rạc","Sàn kết nối giao thương nội bộ","Cơ hội kinh doanh"],
       ["6","Mentoring thủ công","Quản lý ghép cặp & hành trình","Mở rộng quy mô"],
       ["7","Hồ sơ sơ sài","Hồ sơ & thẻ hội viên số","Danh thiếp số"],
       ["8","90 ngày đầu","Hành trình chào đón tự động","Gắn bó sớm"],
       ["9","Chưa ghi nhận năng động","Tích điểm & bảng vinh danh","Khích lệ tham gia"],
       ["10","Chi hội & truyền thông","Quản lý chi hội & hub nội dung","Gắn kết toàn Hội"]],
      y=Inches(1.85),fs=10.5)

# 11 Divider II
s=slide(); divider(s,"Phần II","Đi sâu từng giải pháp",
    "Mỗi điểm nghẽn: hiện trạng, cách giải, mô phỏng giao diện thật, kết quả.")

# 12 Issue 1
s=slide(); head(s,"Điểm nghẽn 01","Dữ liệu hội viên","đang nằm rải rác nhiều nơi",kc=AMBER,emc=AMBER)
two_panel(s,("Hiện tại",["Danh sách ở nhiều file Excel khác nhau","Lịch sử tương tác trôi trên mạng xã hội",
    "Bàn giao nhiệm kỳ dễ thất lạc","Khó trả lời nhanh hội viên đã tham gia gì"]),
    ("Khi có nền tảng",["Một kho dữ liệu hội viên duy nhất","Mọi hoạt động được ghi lại tự động",
    "Bàn giao nhiệm kỳ trọn vẹn dữ liệu","Tra cứu hồ sơ 360° trong vài giây"]))

# 13 Profile 360
s=slide(); head(s,"Giải pháp 01 · Mô phỏng giao diện","Hồ sơ hội viên","360°",kc=GREEN)
txt(s,Inches(0.7),Inches(1.8),Inches(11.9),Inches(0.45),
    [[("Nguyễn Thị An  ·  CEO An Phát  ·  Hạng Tỏa Sáng  ·  Hội viên năng động  ·  2.480 điểm",13,NAVY,True)]])
table(s,["Widget","Nội dung"],
      [["Mức độ gắn kết","88 / 100 (biểu đồ tròn)"],
       ["Hoạt động 6 tháng","Biểu đồ cột theo tháng"],
       ["Hạn gia hạn","02/2027, còn hiệu lực"],
       ["Sự kiện đã dự","12 (năm nay 5)"],
       ["Kết nối giao thương","8, đang trao đổi 2"],
       ["Vai trò Mentoring","Mentor, đang dẫn 3 mentee"],
       ["Dòng thời gian","Diễn giả, gia hạn, ghép mentee..."],
       ["Quyền lợi theo hạng","Vé ưu tiên, 8 ưu đãi, gian hàng"]],y=Inches(2.3),w=Inches(11.0),fs=11)
src(s,"Mô phỏng minh họa kiểu hồ sơ khách hàng (CRM) với hơn 6 widget gồm biểu đồ.")

# 14 Issue 2
s=slide(); head(s,"Điểm nghẽn 02","Gia hạn thủ công khiến","hội viên dễ rời tổ chức vì lãng quên",kc=AMBER,emc=AMBER)
bullets(s,["Văn phòng Hội rà Excel, nhắn từng người khi tới hạn",
           "Hội viên bận, lỡ nhắc rồi vô tình rời tổ chức",
           "Chưa có cảnh báo sớm ai sắp hết hạn, ai cần quan tâm",
           "Mỗi nhiệm kỳ lặp lại bài toán cũ, không tích lũy"],y=Inches(1.9),w=Inches(6.0))
flow(s,[("Trước 30 ngày","Phát hiện sắp hết hạn","Hệ thống tự rà soát",SOFT),
        ("Trước 30/7/1 ngày","Tự nhắc Zalo & email","Gia hạn 1 chạm",RGBColor(0xEA,0xF2,0xFF)),
        ("Cần quan tâm","Báo Ban Chấp hành","Chăm sóc tận tình",RGBColor(0xEE,0xF8,0xF3))],y=Inches(4.1),h=Inches(1.7))
quote(s,"Phần lớn việc rời tổ chức chỉ vì quên gia hạn. Điều này hoàn toàn ngăn được.",y=Inches(2.0),w=Inches(6.0))

# 15 Solution 2
s=slide(); head(s,"Giải pháp 02 · Mô phỏng giao diện","Gia hạn tự động","và lời nhắc ấm áp qua Zalo",kc=GREEN)
bullets(s,["Bảng theo dõi gia hạn: sắp đến hạn, đã tự gia hạn, cần quan tâm",
           "Biểu đồ tỷ lệ tiếp tục đồng hành theo tháng",
           "65% hội viên tự gia hạn online, không cần nhắc tay"],y=Inches(1.95),w=Inches(6.0))
box(s,Inches(7.0),Inches(1.9),Inches(5.6),Inches(4.4),RGBColor(0xDF,0xE7,0xEF),LINE)
box(s,Inches(7.0),Inches(1.9),Inches(5.6),Inches(0.55),RGBColor(0x00,0x68,0xFF))
txt(s,Inches(7.2),Inches(1.97),Inches(5.2),Inches(0.4),[[("HAWEE · Zalo",12,WHITE,True)]])
txt(s,Inches(7.25),Inches(2.65),Inches(5.1),Inches(3.5),
    [[("HAWEE thân gửi chị Mai Phương 💐",12,NAVY,True)],
     [("Quyền lợi hội viên hạng Tỏa Sáng của chị sẽ làm mới vào 25/06/2026.",11.5,INK,False)],
     [(" ",6,INK,False)],
     [("Tiếp tục đồng hành chỉ với 1 phút, giữ trọn quyền lợi & ưu đãi đối tác.",11.5,INK,False)],
     [("[ Gia hạn ngay › ]",12,ACCENT,True)]],sp=1.15)
src(s,"Mô phỏng minh họa giao diện vận hành & tin nhắn Zalo tự động.")

# 16 Impact 2
s=slide(); head(s,"Kết quả 02","Giữ chân vững hơn","là doanh thu & mạng lưới nhân lên",kc=GREEN)
kpis(s,[("+12 điểm %","tỷ lệ hội viên tiếp tục đồng hành",GREEN),
        ("+1,2 tỷ","VND/năm hội phí giữ lại (minh họa)",NAVY),
        ("-80%","thời gian Văn phòng Hội xử lý gia hạn",ACCENT)],y=Inches(2.3))
quote(s,"Khi không ai rời tổ chức chỉ vì lãng quên, mỗi nhiệm kỳ Hội khởi đầu trên nền tảng vững hơn.")
src(s,"Mô hình minh họa. Thay giả định bằng số thực của HAWEE khi khảo sát.")

# 17 Issue 3
s=slide(); head(s,"Điểm nghẽn 03","Gắn kết chỉ","rộn ràng quanh sự kiện rồi lắng xuống",kc=AMBER,emc=AMBER)
two_panel(s,("Hiện tại",["Đăng ký qua form rời rạc hoặc nhắn tay","Điểm danh bằng giấy, không khớp dữ liệu",
    "Sau sự kiện thiếu kết nối lại","Giữa hai kỳ, hội viên ít chạm tới Hội"]),
    ("Khi có nền tảng",["Tạo sự kiện, bán vé theo hạng","Quét QR điểm danh, gắn vào hồ sơ",
    "Tự động cảm ơn & khảo sát sau sự kiện","Bảng tin cộng đồng giữ kết nối liên tục"]))

# 18 Solution 3
s=slide(); head(s,"Giải pháp 03 · Mô phỏng giao diện","Sự kiện, quét QR","và bảng tin cộng đồng",kc=GREEN)
bullets(s,[[("Trang sự kiện: ",13,NAVY,True),("142 đăng ký, 98 đã thanh toán vé, điểm danh QR; vé theo 4 hạng.",13,INK,False)],
           [("Bảng tin cộng đồng: ",13,NAVY,True),("gương mặt hội viên, nhu cầu giao thương, sự kiện sắp tới.",13,INK,False)],
           [("Giữ tương tác liên tục ",13,NAVY,True),("giữa các kỳ sự kiện, không để cộng đồng nguội.",13,INK,False)]],
        y=Inches(2.1),w=Inches(11.9))
quote(s,"Sự kiện không kết thúc khi tan tiệc, nó tiếp tục sống trong cộng đồng.")
src(s,"Mô phỏng minh họa trang sự kiện & bảng tin cộng đồng.")

# 19 QR
s=slide(); head(s,"Trải nghiệm tại sự kiện","Quét QR","là vào ngay, không xếp hàng",kc=GREEN)
qr=box(s,Inches(1.0),Inches(2.1),Inches(3.0),Inches(3.6),NAVY)
box(s,Inches(1.4),Inches(2.5),Inches(2.2),Inches(2.2),WHITE,ACCENT)
txt(s,Inches(1.4),Inches(3.2),Inches(2.2),Inches(0.6),[[("[ QR ]",18,NAVY,True)]],align=PP_ALIGN.CENTER)
box(s,Inches(2.05),Inches(5.0),Inches(0.5),Inches(0.5),GREEN,round_=True)
txt(s,Inches(2.05),Inches(5.05),Inches(0.5),Inches(0.4),[[("✓",16,WHITE,True)]],align=PP_ALIGN.CENTER)
bullets(s,["Hội viên đưa QR trên điện thoại, điểm danh trong 2 giây",
           "Tự động ghi vào hồ sơ 360° và cộng điểm hoạt động",
           "Ban tổ chức thấy ai đã đến theo thời gian thực",
           "Không sổ giấy, không nhầm lẫn danh sách"],x=Inches(4.6),y=Inches(2.3),w=Inches(8.0))
src(s,"Minh họa khái niệm trải nghiệm quét QR điểm danh tại sự kiện.")

# 20 Issue 4
s=slide(); head(s,"Điểm nghẽn 04","Ban Chấp hành","chưa có bức tranh số để nhìn toàn cảnh",kc=AMBER,emc=AMBER)
bullets(s,["Câu hỏi “hội viên nhận được gì từ hội phí” đang trả lời bằng cảm nhận",
           "Khó biết sớm ai đang xa dần để chăm sóc kịp thời",
           "Khó thuyết phục nhà tài trợ vì thiếu số liệu",
           "Quyết định nhiệm kỳ dựa nhiều vào trực giác"],y=Inches(2.0),w=Inches(11.9))
quote(s,"Khi nhìn được bằng số, Ban Chấp hành chủ động dẫn dắt thay vì phản ứng theo sự việc.")

# 21 Dashboard health
s=slide(); head(s,"Giải pháp 04 · Mô phỏng giao diện","Bảng điều hành","Sức khỏe Hội",kc=GREEN)
add_chart(s,XL_CHART_TYPE.PIE,["Cơ Bản","Hội Tụ","Tỏa Sáng","Lan Tỏa"],
          [("Hạng",[22,28,28,22])],Inches(0.7),Inches(1.85),Inches(3.9),Inches(2.5),"Cơ cấu theo hạng",legend=True)
add_chart(s,XL_CHART_TYPE.COLUMN_CLUSTERED,["Cơ Bản","Hội Tụ","Tỏa Sáng","Lan Tỏa"],
          [("Gắn kết",[55,70,86,95])],Inches(4.8),Inches(1.85),Inches(3.9),Inches(2.5),"Mức gắn kết theo hạng",legend=False)
add_chart(s,XL_CHART_TYPE.LINE,["T1","T3","T6","T9","T12"],
          [("Hội viên",[470,484,492,503,512])],Inches(8.9),Inches(1.85),Inches(3.7),Inches(2.5),"Tăng trưởng hội viên",legend=False)
table(s,["Hội viên","Tín hiệu","Gợi ý"],
      [["Phạm Anh T.","Chưa dự sự kiện 6 tháng","Mời cà phê kết nối"],
       ["Vũ Hồng L.","Chưa mở 3 thông báo","Gọi hỏi thăm"]],y=Inches(4.6),w=Inches(11.9),fs=11)
src(s,"Mô phỏng minh họa: 9+ widget gồm biểu đồ tròn, cột, đường xu hướng, bảng chăm sóc sớm.")

# 22 Issue 5
s=slide(); head(s,"Điểm nghẽn 05 · trọng tâm nhiệm kỳ III","Giao thương","còn dựa nhiều vào quan hệ cá nhân",kc=AMBER,emc=AMBER)
two_panel(s,("Hiện tại",["Kết nối chủ yếu qua gặp gỡ tại sự kiện","Chưa có nơi đăng nhu cầu hay chào năng lực",
    "Hội chưa ghi nhận được giá trị giao thương","Định hướng giao thương thiếu công cụ"]),
    ("Khi có nền tảng",["Sàn kết nối, tìm theo ngành & nhu cầu","Đăng nhu cầu, gửi kết nối 1 chạm",
    "Gợi ý đối tác phù hợp tự động","Ghi nhận kết nối & câu chuyện thành công"]))

# 23 B2B
s=slide(); head(s,"Giải pháp 05 · Mô phỏng giao diện","Sàn kết nối","Giao thương nội bộ",kc=GREEN)
bullets(s,["Bộ lọc theo ngành: F&B, Bán lẻ, Logistics, Mỹ phẩm, Bất động sản",
           "214 doanh nghiệp hội viên · 37 nhu cầu mở · 28 kết nối quý 2"],y=Inches(1.85),w=Inches(11.9))
table(s,["Gợi ý kết nối thông minh","Mức khớp nhu cầu"],
      [["An Phát (F&B) ✕ Bao bì Minh Long","92%"],
       ["Mỹ phẩm Hạ ✕ Chuỗi bán lẻ Vy","87%"],
       ["Logistics P. ✕ An Phát","81%"]],y=Inches(3.0),w=Inches(8.5),fs=12)
quote(s,"“Qua sàn HAWEE, chúng tôi chốt hợp đồng phân phối trong 3 tuần.”",y=Inches(5.5),w=Inches(11.9))
src(s,"Mô phỏng minh họa: bộ lọc ngành, thẻ doanh nghiệp, gợi ý thông minh, câu chuyện thành công.")

# 24 Impact 5
s=slide(); head(s,"Kết quả 05","Từ “hội networking” thành","nơi tạo ra cơ hội kinh doanh",kc=GREEN)
kpis(s,[("Mọi lúc","kết nối không chờ tới kỳ sự kiện",ACCENT),
        ("Đo được","số kết nối & giá trị giao thương",GREEN),
        ("Lý do ở lại","hội viên thấy lợi ích kinh doanh",GOLD)],y=Inches(2.3))
quote(s,"Khi hội viên tìm được đối tác và đơn hàng từ HAWEE, câu hỏi hội phí để làm gì tự có câu trả lời.")

# 25 Mentoring
s=slide(); head(s,"Điểm nghẽn 06 › Giải pháp","Số hóa “viên ngọc”","HAWEE Mentoring")
bullets(s,["Hiện tại: ghép cặp, theo dõi 12 tháng, nhắc lịch còn thủ công",
           "Giải pháp: ghép cặp theo ngành, theo dõi hành trình, nhắc lịch tự động, thu phản hồi"],
        y=Inches(1.9),w=Inches(6.0))
table(s,["Mentor","Mentee","Ngành","Tiến độ"],
      [["Nguyễn T. An","Lê M. Chi","F&B","8/12"],
       ["Trần B. Hà","Vũ K. Linh","Bán lẻ","3/12"]],y=Inches(1.9),x=Inches(7.0),w=Inches(5.6),fs=11)
quote(s,"Mở rộng chương trình điểm mà không tăng tải vận hành cho Ban Mentoring.",y=Inches(4.0),w=Inches(6.0))
src(s,"Mô phỏng minh họa module Mentoring · 48 cặp đang hoạt động, 92% buổi đúng lịch.")

# 26 Digital card
s=slide(); head(s,"Điểm nghẽn 07 › Giải pháp","Hồ sơ doanh nhân &","thẻ hội viên số")
bullets(s,["Hội viên chưa có “danh thiếp số” để thể hiện năng lực & tìm thấy nhau",
           "Quyền lợi 4 hạng chưa hiển thị rõ ràng",
           "Giải pháp: hồ sơ doanh nhân + thẻ số QR theo hạng + ưu đãi đối tác"],y=Inches(2.0),w=Inches(7.0))
box(s,Inches(8.2),Inches(2.1),Inches(4.0),Inches(3.4),NAVY2)
txt(s,Inches(8.45),Inches(2.4),Inches(3.5),Inches(2.8),
    [[("HAWEE · HẠNG TỎA SÁNG",10,RGBColor(0x9E,0xC7,0xE8),True)],
     [("Nguyễn Thị An",16,WHITE,True)],
     [("CEO · An Phát F&B · từ 2019",10,RGBColor(0xCF,0xDD,0xED),False)],
     [(" ",10,WHITE,False)],
     [("[ QR ]  Quét để điểm danh & mở quyền lợi",10,RGBColor(0x9E,0xC7,0xE8),False)]],sp=1.3)
src(s,"Mô phỏng minh họa thẻ hội viên số theo 4 hạng.")

# 27 Issue 8 onboarding
s=slide(); head(s,"Điểm nghẽn 08 (mở rộng) › Giải pháp","90 ngày đầu","quyết định hội viên có gắn bó lâu dài")
bullets(s,["Hội viên mới không gắn kết sớm thì rất dễ rời tổ chức về sau",
           "Hiện việc chào đón còn tùy hứng, chưa thành quy trình",
           "Giải pháp: hành trình chào đón tự động 90 ngày, gợi ý sự kiện & kết nối phù hợp"],
        y=Inches(1.9),w=Inches(6.0))
flow(s,[("Ngày 1","Lời chào & cẩm nang quyền lợi","Gửi qua Zalo",SOFT),
        ("Tuần 2 đến 4","Mời sự kiện đầu tiên","Gợi ý 3 hội viên nên kết nối",RGBColor(0xEA,0xF2,0xFF)),
        ("Ngày 90","Điểm lại hành trình","Mời tham gia chuyên sâu",RGBColor(0xEE,0xF8,0xF3))],y=Inches(4.0),h=Inches(1.8))
src(s,"Cơ sở: hội viên không gắn kết trong 90 ngày đầu rời tổ chức nhiều hơn đáng kể.")

# 28 Gamification
s=slide(); head(s,"Điểm nghẽn 09 (mở rộng) › Giải pháp","Tích điểm","và vinh danh hội viên năng động",kc=GREEN)
bullets(s,["Dự sự kiện, làm mentor, giới thiệu hội viên mới, kết nối giao thương đều được cộng điểm",
           "Hội nhìn rõ ai là hạt nhân năng động để ghi nhận & mời dẫn dắt",
           "Hạng hoạt động: Kim Cương, Vàng, Bạc, kèm quyền lợi & vinh danh tại đại hội"],
        y=Inches(1.85),w=Inches(5.4))
table(s,["#","Hội viên","Đóng góp","Điểm"],
      [["1","Lê Thanh Hà","Mentor · 9 sự kiện · 5 kết nối","3.120"],
       ["2","Trần Mỹ","8 sự kiện · 4 giới thiệu","2.760"],
       ["3","Nguyễn Thị An","Mentor · 7 sự kiện","2.480"],
       ["4","Vũ Kim","6 sự kiện · 3 kết nối","2.050"],
       ["5","Phạm Hoa","5 sự kiện · 2 giới thiệu","1.890"]],y=Inches(1.85),x=Inches(6.4),w=Inches(6.2),fs=10.5)
src(s,"Mô phỏng minh họa hệ thống tích điểm & bảng vinh danh hội viên năng động.")

# 29 Chapters
s=slide(); head(s,"Điểm nghẽn 10 (mở rộng) › Giải pháp","Gắn kết","cả chi hội & ban chuyên môn")
cards(s,[("Chi hội & ban","Trang riêng","CLB Phong cách, ban Mentoring, ban Giao thương: lịch, thành viên, hoạt động riêng."),
         ("Nội dung","Hub tập trung","Bản tin, gương mặt hội viên, thư viện tài liệu, đo lượt xem & quan tâm."),
         ("Khảo sát","Lấy ý kiến","Khảo sát, biểu quyết, đo mức hài lòng theo từng nhóm."),
         ("Tổng hợp","Một bức tranh","Mọi hoạt động chi hội đổ về bảng điều hành chung của Hội.")],y=Inches(2.1),h=Inches(2.7))
quote(s,"Mỗi chi hội có không gian riêng, nhưng tất cả cùng một mái nhà dữ liệu.")

# 30 Divider III
s=slide(); divider(s,"Phần III","Tự động hóa, AI & Website",
    "Những trải nghiệm hội viên chạm vào mỗi ngày.")

# 31 Automation journey
s=slide(); head(s,"Trụ cột tự động hóa","Một hành trình chăm sóc","tự động qua Zalo & email",kc=GREEN)
flow(s,[("Gia nhập","Chào đón","Onboarding & cẩm nang quyền lợi",SOFT),
        ("Trong kỳ","Gắn kết","Mời sự kiện đúng mối quan tâm",SOFT),
        ("Trước hạn","Nhắc gia hạn","Tự động qua Zalo & email",RGBColor(0xEA,0xF2,0xFF)),
        ("Cần quan tâm","Chăm sóc","Gợi ý BCH quan tâm kịp thời",RGBColor(0xEE,0xF8,0xF3))],y=Inches(2.3),h=Inches(2.0))
quote(s,"Zalo là kênh người Việt mở tin nhiều & tin tưởng nhất. Đây là lợi thế riêng của HAWEE.")
src(s,"Nguồn: Zalo có hơn 75 triệu người dùng/tháng tại Việt Nam (Adsota 2025).")

# 32 AI verify new member
s=slide(); head(s,"Trụ cột trợ lý ảo · Khi có hội viên mới","Xác minh hội viên mới","bằng AI, duyệt nhanh",kc=GREEN)
table(s,["Trường thông tin","AI tự tổng hợp"],
      [["Doanh nghiệp","Công ty TNHH Mỹ phẩm Hạ Vy"],
       ["Mã số thuế / Ngành","0312•••456 · Mỹ phẩm"],
       ["Thành lập / Quy mô","2017 · 50 đến 100 nhân sự"],
       ["Độ tin cậy","92% · 5 nguồn đối chiếu"],
       ["Khớp tiêu chí hội viên","Đạt · đề xuất hạng Hội Tụ"]],y=Inches(1.9),x=Inches(0.7),w=Inches(6.2),fs=11)
bullets(s,["Thay vì tìm kiếm & tổng hợp thủ công, AI lập sẵn báo cáo doanh nghiệp",
           "Đối chiếu nhiều nguồn công khai: pháp nhân, ngành, quy mô, truyền thông",
           "Ban Chấp hành chỉ xem & bấm duyệt, rút ngắn thời gian kết nạp",
           "Con người vẫn ra quyết định cuối cùng"],x=Inches(7.1),y=Inches(2.0),w=Inches(5.5),sz=12.5)
src(s,"Mô phỏng minh họa. AI tổng hợp từ nguồn công khai; quyết định kết nạp thuộc Ban Chấp hành.")

# 33 AI assistant
s=slide(); head(s,"Trụ cột trợ lý ảo","Trợ lý ảo","đồng hành cùng hội viên 24/7",kc=GREEN)
box(s,Inches(0.7),Inches(1.9),Inches(6.6),Inches(3.6),SOFT,LINE)
txt(s,Inches(0.95),Inches(2.1),Inches(6.1),Inches(3.2),
    [[("Hội viên: ",12,NAVY,True),("HAWEE Connect Q3 khi nào? Hạng Tỏa Sáng có vé miễn phí?",12,INK,False)],
     [(" ",6,INK,False)],
     [("Trợ lý HAWEE: ",12,ACCENT,True),("18/07/2026, 18:00 tại KS Rex. Hạng Tỏa Sáng được vé miễn phí + bàn ưu tiên. Em giữ chỗ nhé?",12,INK,False)],
     [(" ",6,INK,False)],
     [("Hội viên: ",12,NAVY,True),("Giữ giúp mình. Cho xem hạn gia hạn luôn.",12,INK,False)],
     [(" ",6,INK,False)],
     [("Trợ lý HAWEE: ",12,ACCENT,True),("Đã giữ chỗ. Quyền lợi làm mới 25/06/2026, em gửi liên kết qua Zalo nhé.",12,INK,False)]],sp=1.15)
bullets(s,["Trả lời tức thì: sự kiện, quyền lợi theo hạng, hạn gia hạn",
           "Chuyển đúng người: Văn phòng Hội, Ban Giao thương, Mentoring",
           "Giảm tải hỏi đáp lặp lại cho Văn phòng Hội",
           "Định vị đúng: trợ lý & điều hướng, không thay con người"],x=Inches(7.6),y=Inches(2.1),w=Inches(5.0),sz=12)
src(s,"Mô phỏng minh họa hội thoại. Triển khai ở Giai đoạn 3.")

# 34 Community + newsletter
s=slide(); head(s,"Cộng đồng & bản tin","Trang cộng đồng sôi nổi","và bản tin định kỳ tự động",kc=GREEN)
bullets(s,[[("Có ",13,GREEN,True),("trang cộng đồng riêng: bảng tin, nhóm chi hội, gương mặt hội viên, hỏi đáp, tương tác (cảm xúc, bình luận, chia sẻ).",13,INK,False)],
           [("Có ",13,GREEN,True),("bản tin định kỳ tự động tổng hợp hoạt động cộng đồng, gửi hằng tháng qua email & Zalo.",13,INK,False)]],
        y=Inches(1.95),w=Inches(11.9))
table(s,["Bản tin HAWEE · Tháng 7/2026","Nội dung tự tổng hợp"],
      [["Sự kiện sắp tới","HAWEE Connect Q3 · 18/07"],
       ["Gương mặt hội viên","Câu chuyện chị Nguyễn Thị An"],
       ["Cơ hội giao thương","12 nhu cầu mới chờ kết nối"],
       ["Hội viên mới","Chào mừng 8 hội viên"]],y=Inches(3.3),w=Inches(9.0),fs=11)
src(s,"Mô phỏng minh họa trang cộng đồng & bản tin tự động.")

# 35 Website vibrant
s=slide(); head(s,"Bộ mặt đối ngoại","Website chính thức","sôi nổi như một cộng đồng sống",kc=GREEN)
box(s,Inches(0.7),Inches(1.85),Inches(11.9),Inches(2.1),NAVY2)
txt(s,Inches(1.0),Inches(2.05),Inches(8.0),Inches(1.7),
    [[("● 18 hội viên đang trực tuyến",11,RGBColor(0x9E,0xC7,0xE8),True)],
     [("Hội tụ. Kết nối. Vươn tầm.",24,WHITE,True)],
     [("Cộng đồng nữ doanh nhân TP.HCM, nơi mỗi kết nối mở ra một cơ hội.",12,RGBColor(0xCF,0xDD,0xED),False)],
     [("[ Trở thành hội viên › ]   [ Khám phá cộng đồng ]",12,GOLD,True)]],sp=1.2)
bullets(s,["Hero sống động, chip hội viên đang trực tuyến, cụm avatar hội viên (+500)",
           "Hàng thẻ “Đang sôi nổi trong cộng đồng”: sự kiện, gương mặt hội viên, cơ hội giao thương",
           "Ô đăng ký bản tin ngay trên trang chủ",
           "4 chỉ số nổi bật: 500+ hội viên, 200+ chương trình, 14.000+ lượt dự, 105 tỷ vì cộng đồng"],
        y=Inches(4.2),w=Inches(11.9),sz=12.5)
src(s,"Mô phỏng minh họa trang chủ website đối ngoại sôi nổi của HAWEE.")

# 36 Analytics dashboard
s=slide(); head(s,"Sau khi có website · Báo cáo & phân tích","Bảng điều hành","Báo cáo & Phân tích chuyên sâu",kc=GREEN)
add_chart(s,XL_CHART_TYPE.LINE,["T1","T3","T6","T9","T12"],
          [("Hội viên",[470,484,492,503,512]),("Giữ chân %",[72,74,76,79,81])],
          Inches(0.7),Inches(1.85),Inches(4.0),Inches(2.5),"Tăng trưởng & Giữ chân",legend=True)
add_chart(s,XL_CHART_TYPE.COLUMN_STACKED,["Q1","Q2","Q3","Q4"],
          [("Hội phí",[34,40,46,52]),("Sự kiện",[14,18,22,26]),("Tài trợ",[10,12,16,18])],
          Inches(4.85),Inches(1.85),Inches(4.0),Inches(2.5),"Doanh thu theo nguồn",legend=True)
add_chart(s,XL_CHART_TYPE.DOUGHNUT,["Cơ Bản","Hội Tụ","Tỏa Sáng","Lan Tỏa"],
          [("Hạng",[22,28,28,22])],Inches(9.0),Inches(1.85),Inches(3.6),Inches(2.5),"Cơ cấu theo hạng",legend=True)
table(s,["Phễu kết nạp","Cohort giữ chân","Top ngành"],
      [["Quan tâm 240 › Đăng ký 150","Theo tháng gia nhập","F&B, Bán lẻ"],
       ["AI xác minh 132 › Kết nạp 118","Đậm = giữ chân cao","Mỹ phẩm, BĐS"]],y=Inches(4.55),w=Inches(11.9),fs=11)
src(s,"Mô phỏng minh họa: đường, cột chồng, tròn, phễu, cohort, xếp hạng. Lọc & xuất báo cáo nhiều tình huống.")

# 37 Sponsor ROI
s=slide(); head(s,"Đòn bẩy ngân sách","Quản lý &","đo hiệu quả tài trợ",kc=GREEN)
table(s,["Quyền lợi","Cam kết","Thực hiện"],
      [["Logo cổng sự kiện","7","7"],
       ["Bài viết cộng đồng","4","4"],
       ["Gian hàng giao thương","2","1/2"]],y=Inches(1.9),x=Inches(0.7),w=Inches(6.0),fs=12)
bullets(s,["Theo dõi cam kết & quyền lợi từng nhà tài trợ",
           "Báo cáo tiếp cận & khách quan tâm, nhà tài trợ thấy rõ hiệu quả",
           "Dễ tái ký & nâng gói, tăng nguồn lực cho Hội"],x=Inches(7.0),y=Inches(2.0),w=Inches(5.6))
quote(s,"Tài trợ là nguồn lực quan trọng. Đo được hiệu quả là cách giữ & mở rộng nguồn lực ấy.")
src(s,"Mô phỏng minh họa báo cáo tài trợ: 7 sự kiện, 2.100 lượt tiếp cận, 34 khách quan tâm.")

# 38 Divider IV
s=slide(); divider(s,"Phần IV","Tổng hợp & lộ trình",
    "Từ bức tranh tính năng đến kế hoạch có thể ra quyết định.")

# 39 Feature map
s=slide(); head(s,"Bản đồ tính năng đầy đủ","Tất cả gói gọn trong","một nền tảng")
cards(s,[("Nền móng","Dữ liệu & Hội viên","Kho dữ liệu tập trung · Hồ sơ 360° · 4 hạng · Gia hạn tự động · Nhắc qua Zalo"),
         ("Giá trị","Gắn kết & Giao thương","Sự kiện + QR · Sàn giao thương · Mentoring · Tích điểm & vinh danh · Cộng đồng · Thẻ số"),
         ("Vươn tầm","Thông minh & Mở rộng","Bảng điều hành · Chăm sóc sớm · Trợ lý ảo · Xác minh AI · Hiệu quả tài trợ")],
      y=Inches(2.0),h=Inches(3.0))
quote(s,"7 trụ cột HAWEE đề ra, cùng các tính năng mở rộng, trên một nền dữ liệu duy nhất.")

# 40 Roadmap
s=slide(); head(s,"Lộ trình triển khai","Làm chắc gốc trước,","không ôm tất cả cùng lúc")
cards(s,[("Giai đoạn 1 · Nền tảng","Một nguồn duy nhất","Website & cổng hội viên · Kho dữ liệu & Hồ sơ 360° · 4 hạng & gia hạn tự động · Nhắc Zalo. › Thắng nhanh, đo được"),
         ("Giai đoạn 2 · Giá trị","Kết nối tạo giá trị","Sự kiện + QR + cộng đồng · Sàn giao thương · Mentoring · Tích điểm & vinh danh. › Gắn kết liên tục"),
         ("Giai đoạn 3 · Vươn tầm","Dữ liệu dẫn dắt","Bảng điều hành & chăm sóc sớm · Trợ lý ảo · Hiệu quả tài trợ · Ứng dụng & Zalo Mini App. › Tăng giữ chân & tài trợ")],
      y=Inches(2.0),h=Inches(3.2))

# 41 Phase matrix
s=slide(); head(s,"Ma trận tính năng theo giai đoạn","Cái gì làm","trước, cái gì làm sau")
table(s,["Nhóm tính năng","GĐ 1","GĐ 2","GĐ 3"],
      [["Kho dữ liệu & Hồ sơ 360°","●","",""],
       ["4 hạng & gia hạn tự động","●","",""],
       ["Nhắc qua Zalo & email","●","",""],
       ["Sự kiện + QR + cộng đồng","","●",""],
       ["Sàn giao thương","","●",""],
       ["Mentoring & Tích điểm vinh danh","","●",""],
       ["Bảng điều hành & chăm sóc sớm","","","●"],
       ["Trợ lý ảo · Xác minh AI · Tài trợ · Ứng dụng","","","●"]],y=Inches(1.9),fs=11.5)

# 42 Value
s=slide(); head(s,"Giá trị mang lại","Mỗi bên đều","được lợi",kc=GREEN)
cards(s,[("Hội","Vững & minh bạch","Giữ chân tốt hơn, dữ liệu chứng minh giá trị, tăng tài trợ"),
         ("Hội viên","Lợi ích thật","Tìm đúng đối tác, quyền lợi rõ ràng, ít thao tác"),
         ("Ban Chấp hành","Quyết định bằng số","Bảng điều hành thay cảm tính, hành động kịp thời"),
         ("Đối tác","Thấy hiệu quả","Dễ tái ký & nâng gói tài trợ")],y=Inches(2.0),h=Inches(2.7))
quote(s,"Tầm nhìn của HAWEE đã rõ. Nền tảng này biến tầm nhìn ấy thành vận hành mỗi ngày.")

# 43 Why build
s=slide(); head(s,"Vì sao xây riêng cho HAWEE","Một tài sản số","thuộc về Hội, lớn lên cùng Hội",kc=GREEN)
cards(s,[("Chủ quyền","Dữ liệu của HAWEE","Thông tin hội viên trong tầm kiểm soát của Hội, không phụ thuộc bên thứ ba."),
         ("Bản sắc","May đo trọn vẹn","Đúng 4 hạng, đúng Mentoring, đúng định hướng giao thương."),
         ("Người Việt","Gắn liền Zalo","Nhắc nhở & chăm sóc qua Zalo, thanh toán nội địa quen thuộc."),
         ("Bền vững","Chi phí hợp lý dài hạn","Đầu tư một lần & bảo trì hợp lý, không tăng phí theo từng hội viên.")],
      y=Inches(2.0),h=Inches(2.8))
quote(s,"Với HAWEE, vận hành hội viên là lợi thế cần sở hữu, không phải thứ đi thuê.")

# 44 Risks
s=slide(); head(s,"Lường trước để đi chắc","Mọi rủi ro","đều có cách hóa giải")
table(s,["Điều cần lưu tâm","Cách hóa giải"],
      [["Hội viên bận, ngại công cụ mới","Giao diện tối giản, nhắc qua Zalo, không bắt học hệ thống"],
       ["Dữ liệu nhạy cảm của nữ lãnh đạo","Hội sở hữu dữ liệu, phân quyền chặt, tuân thủ bảo vệ dữ liệu cá nhân"],
       ["Lo làm xong rồi để đó","Bàn giao kèm quy trình vận hành & người phụ trách cộng đồng"],
       ["Triển khai lớn dễ rủi ro","Lộ trình 3 giai đoạn, Giai đoạn 1 nhỏ, chắc, đo được trước"],
       ["Kỳ vọng trợ lý ảo quá cao","Định vị đúng là trợ lý & điều hướng, đặt kỳ vọng phù hợp"]],y=Inches(1.95),fs=12)

# 45 Next steps
s=slide(); head(s,"Bước tiếp theo","Khởi đầu","nhỏ, chắc, đo được",kc=GREEN)
flow(s,[("Bước 1","Khảo sát nhu cầu","Trao đổi Văn phòng Hội & hội viên, lấy số thực",SOFT),
        ("Bước 2","Chốt phạm vi Giai đoạn 1","Tính năng nền tảng & mốc thắng nhanh",SOFT),
        ("Bước 3","Lộ trình & đồng hành","Kế hoạch 3 giai đoạn & hợp tác dài hạn",RGBColor(0xEE,0xF8,0xF3))],
     y=Inches(2.2),h=Inches(1.9))
quote(s,"HAWEE đã có tầm nhìn. Hãy bắt đầu từ Giai đoạn 1, để tầm nhìn ấy thành hiện thực mỗi ngày.")
src(s,"Có thể kèm bộ câu hỏi khảo sát nhu cầu hội viên & Ban Chấp hành theo yêu cầu.")

# 46 End
s=slide(); cover(s,"HAWEE 2025 đến 2030","Hội tụ. Kết nối.\nVươn tầm.",
    "Một nền tảng xây riêng cho HAWEE, vận hành tầm nhìn ấy mỗi ngày.",
    "HAWEE      Cảm ơn & sẵn sàng trao đổi      2026")

import os
out=os.path.join(os.path.dirname(os.path.abspath(__file__)),"hawee-pitch-deck.pptx")
prs.save(out)
print("Đã tạo", out, "·", len(prs.slides._sldIdLst), "slide")
