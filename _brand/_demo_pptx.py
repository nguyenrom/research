# -*- coding: utf-8 -*-
"""Demo PPTX. Run: .venv-pptx/bin/python _brand/_demo_pptx.py"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptx.util import Inches
from brandkit import pptx as bc, PRIMARY, ACCENT, NEUTRAL

prs = Presentation()
prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)  # 16:9
bc.title_slide(prs, "Blue Coral Pitch", subtitle="Triển khai Dữ liệu & Marketing")

# a content slide proving body text helper
s = prs.slides.add_slide(prs.slide_layouts[6])
bc._text(s, "Nội dung", Inches(0.6), Inches(0.5), Inches(8), Inches(0.8), size=28, bold=True, color=PRIMARY["700"])
bc._text(s, "Slide nội dung dùng màu Trust Blue + Coral từ token.", Inches(0.6), Inches(1.6), Inches(11), Inches(1), size=18, color=NEUTRAL["ink"])

out = os.path.join(os.path.dirname(__file__), "templates", "demo-deck.pptx")
os.makedirs(os.path.dirname(out), exist_ok=True)
prs.save(out)
print("OK PPTX ->", out)
