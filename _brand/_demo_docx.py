# -*- coding: utf-8 -*-
"""Demo Word doc. Run: .venv-docx/bin/python _brand/_demo_docx.py"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from brandkit import docx as bc

doc = Document()
bc.apply_base(doc)
bc.cover_title(doc, "Đề xuất triển khai", subtitle="Tài liệu chuẩn bị cho buổi trao đổi", date="10/06/2026")
doc.add_heading("I. Bối cảnh", level=1)
doc.add_paragraph("Đoạn văn theo chuẩn Blue Coral, font và màu lấy từ brand-tokens.json.")
doc.add_heading("II. Phạm vi & chi phí", level=2)
bc.brand_table(doc, ["Hạng mục", "Ngày công", "Chi phí (VND)"],
               [["Nền tảng", "13", "32.000.000"],
                ["Triển khai", "12", "30.000.000"],
                ["Tổng", "25", "62.000.000"]])
bc.footer(doc)

out = os.path.join(os.path.dirname(__file__), "templates", "demo-proposal.docx")
os.makedirs(os.path.dirname(out), exist_ok=True)
doc.save(out)
print("OK Word ->", out)
