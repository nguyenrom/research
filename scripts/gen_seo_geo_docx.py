# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Base font
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

def h1(text):
    p = doc.add_heading(text, level=1)
    return p

def h2(text):
    p = doc.add_heading(text, level=2)
    return p

def para(text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    return p

def bullet(text):
    doc.add_paragraph(text, style="List Bullet")

def numbered(text):
    doc.add_paragraph(text, style="List Number")

# ---- Title ----
title = doc.add_heading("Đề xuất SEO/GEO — Quan điểm và các vấn đề cần làm rõ", level=0)
sub = doc.add_paragraph("Tài liệu nội bộ chuẩn bị cho buổi trao đổi / pitching")
sub.runs[0].italic = True
doc.add_paragraph("Ngày: 02/06/2026")

# ---- I. Nhận định ----
h1("I. Nhận định")

para("1. SEO hiện nay tập trung vào hai phần chính:", bold=True)
bullet("Chiến lược nội dung: viết đúng nhu cầu tìm kiếm của người dùng, đủ sâu để được tin cậy.")
bullet("Cấu trúc website (schema, cách tổ chức trang): để cả Google và các trợ lý AI/LLM (ChatGPT, Gemini, AI Overviews) đọc và hiểu được website. Đây là phần GEO — tối ưu để AI hiểu và đề xuất, không chỉ tối ưu cho Google.")

para("2. Nên chia dự án thành 3 giai đoạn:", bold=True)
numbered("Giai đoạn 1 — Chiến lược: gắn mục tiêu kinh doanh với chiến lược SEO/GEO. Thống nhất đi đâu và đo bằng gì.")
numbered("Giai đoạn 2 — Nội dung và kỹ thuật: thiết kế chiến lược nội dung và chuẩn hóa cấu trúc/schema website.")
numbered("Giai đoạn 3 — Triển khai: thực thi, đo lường và tối ưu.")

para("3. SEO/GEO phải duy trì hàng năm, không phải làm một lần là xong, vì đối thủ cũng làm SEO/GEO.", bold=True)
para("Vì vậy cần cân nhắc thứ tự ưu tiên: SEO/GEO nên đặt trên nền tảng sẵn có, không nên ưu tiên tuyệt đối.")
bullet("Cần làm song song: chương trình gắn kết / giới thiệu qua chính sinh viên, và chiến dịch marketing để tăng nhận diện thương hiệu. Khi người dùng đã gắn hình ảnh thương hiệu với trường, họ sẽ chủ động tìm đến.")
bullet("Ưu tiên cốt lõi vẫn là website: cần chốt nền tảng website trước khi làm SEO/GEO. Tránh trường hợp làm SEO/GEO xong lại sửa lại website, dẫn đến làm lại hai lần (double work), tốn chi phí và thời gian.")

# ---- II. Câu hỏi ----
h1("II. Các vấn đề cần làm rõ")

para("1. User Journey trên website", bold=True)
para("Mọi traffic đều dẫn về website, và mục tiêu cuối là tăng lead/khách hàng, không chỉ tăng lượt truy cập. Cần làm rõ người dùng đi qua những bước nào và đâu là điểm chuyển đổi (đăng ký, để lại thông tin, liên hệ).")

para("2. Cách đo lường và retarget", bold=True)
para("Cần thống nhất cách triển khai tracking (CRM/CDP) và retarget trên website, để đánh giá hành vi người dùng và đo hiệu quả thực sự của cả SEO/GEO lẫn website. Không có lớp đo lường này thì không biết khoản đầu tư có hiệu quả hay không.")

para("3. Phạm vi can thiệp và ràng buộc KPI", bold=True)
para("Cần xác định có phải thay đổi gì ở website hoặc kênh social để đáp ứng SEO/GEO không. Từ đó chốt KPI ràng buộc giữa hai bên và cơ chế đánh giá KPI.")
para("Lưu ý về đo lường GEO: khác với Google, Agent/LLM không có công cụ phân tích để đo hiển thị hay xếp hạng. Ta chỉ tracking thụ động khi người dùng đã vào website. Thậm chí có nghịch lý: nếu nội dung của ta đã đủ trên Agent/LLM, người dùng nhận được câu trả lời ngay tại đó và không vào website. Vì vậy KPI cho GEO cần định nghĩa khác với KPI SEO truyền thống.")

# ---- III. Gợi ý thêm ----
h1("III. Gợi ý thêm (nên chốt trước khi bắt đầu)")
numbered("Chốt số liệu nền (baseline) trước khi làm: traffic hiện tại, từ khóa đang xếp hạng, số lead mỗi tháng. Có mốc thì mới so sánh được kết quả.")
numbered("Làm rõ quyền truy cập website: ai có quyền sửa code/CMS để gắn schema và chỉnh cấu trúc trang.")
numbered("Phân vai rõ ràng: ai làm nội dung, ai làm kỹ thuật, ai làm tracking. Tránh chồng chéo và đùn đẩy.")
numbered("Thống nhất quyền sở hữu dữ liệu: dữ liệu khách/lead thuộc về ai và lưu ở đâu.")
numbered("Tính trước chi phí duy trì hàng năm, vì SEO/GEO phải làm liên tục.")
numbered("Theo dõi đối thủ: họ đang làm gì, để biết mình cần đầu tư đến mức nào.")
numbered("Lường trước rủi ro phụ thuộc nền tảng: Google đổi thuật toán, LLM đổi cách trích dẫn. Cần phương án dự phòng.")

out = "/Users/rom/Projects/research/_bmad-output/planning-artifacts/research/de-xuat-seo-geo-2026-06-02.docx"
doc.save(out)
print("SAVED:", out)
